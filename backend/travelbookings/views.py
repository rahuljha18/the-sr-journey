from rest_framework.decorators import api_view
from rest_framework.response import Response
from rest_framework import status
from django.db.models import Sum
from django.shortcuts import get_object_or_404, render
from datetime import datetime
from django.shortcuts import render, get_object_or_404, redirect
from django.http import HttpResponse
from django.contrib.admin.views.decorators import staff_member_required
from django.views.decorators.http import require_POST

from openpyxl import Workbook

from .models import Booking, Partner, PartnerSettlement


# =========================================================
# BOOKING LIST API
# =========================================================

@api_view(["GET"])
def booking_list(request):

    bookings = Booking.objects.all().order_by("-created_at")

    data = []

    for booking in bookings:

        partner = booking.partner

        data.append({
            "booking_id": booking.booking_id,
            "customer_name": booking.customer_name,
            "mobile": booking.mobile,
            "service_type": booking.service_type,
            "package_name": booking.package_name,
            "travel_date": booking.travel_date,
            "passengers": booking.passengers,
            "cab_category": booking.cab_category,
            "pickup_location": booking.pickup_location,
            "drop_location": booking.drop_location,
            "amount": str(booking.amount),
            "payment_status": booking.payment_status,
            "booking_status": booking.booking_status,

            "partner": (
                partner.partner_name
                if partner
                else None
            ),

            "company_name": (
                partner.company_name
                if partner
                else None
            ),

            "driver_name": (
                partner.driver_name
                if partner
                else None
            ),

            "driver_mobile": (
                partner.driver_mobile
                if partner
                else None
            ),

            "vehicle_number": (
                partner.vehicle_number
                if partner
                else None
            ),

            "partner_cab_category": (
                partner.cab_category
                if partner
                else None
            ),
        })

    return Response(data)


# =========================================================
# CREATE BOOKING API
# =========================================================


from decimal import Decimal, ROUND_HALF_UP

@api_view(["POST"])
def create_booking(request):
    data = request.data

    # Booking amount ko safely Decimal mein convert karein
    try:
        amount = Decimal(str(data.get("amount") or "0"))
        if not amount.is_finite() or amount < 0:
            raise ValueError
    except (ValueError, TypeError, ArithmeticError):
        return Response(
            {"error": "Please provide a valid non-negative booking amount."},
            status=status.HTTP_400_BAD_REQUEST
        )

    # Commission calculation: 5%
    commission_rate = Decimal("5.00")

    commission_amount = (
        amount * commission_rate / Decimal("100")
    ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    partner_payable = amount - commission_amount

    booking = Booking.objects.create(
        booking_id=data.get("booking_id"),
        customer_name=data.get("customer_name"),
        mobile=data.get("mobile"),
        service_type=data.get("service_type"),
        package_name=data.get("package_name", ""),
        travel_date=data.get("travel_date") or None,
        passengers=data.get("passengers", 1),
        cab_category=data.get("cab_category", ""),
        pickup_location=data.get("pickup_location", ""),
        drop_location=data.get("drop_location", ""),
        amount=amount,
        commission_rate=commission_rate,
        commission_amount=commission_amount,
        partner_payable=partner_payable,
        special_requirement=data.get("special_requirement", ""),
    )

    return Response(
        {
            "message": "Booking created successfully",
            "booking_id": booking.booking_id,
            "amount": str(booking.amount),
            "commission_rate": str(booking.commission_rate),
            "commission_amount": str(booking.commission_amount),
            "partner_payable": str(booking.partner_payable),
        },
        status=status.HTTP_201_CREATED
    )



# =========================================================
# BOOKING TRACKING API
# =========================================================

@api_view(["GET"])
def booking_tracking(request, booking_id):

    try:

        booking = Booking.objects.select_related(
            "partner"
        ).get(
            booking_id=booking_id
        )

    except Booking.DoesNotExist:

        return Response(
            {
                "message": "Booking not found"
            },
            status=status.HTTP_404_NOT_FOUND
        )

    partner = booking.partner

    data = {

        "booking_id": booking.booking_id,

        "customer_name": booking.customer_name,

        "service_type": booking.service_type,

        "package_name": booking.package_name,

        "travel_date": booking.travel_date,

        "passengers": booking.passengers,

        "cab_category": booking.cab_category,

        "amount": str(booking.amount),

        "payment_status": booking.payment_status,

        "booking_status": booking.booking_status,

        "partner": (

            {
                "partner_name": partner.partner_name,

                "company_name": partner.company_name,

                "driver_name": partner.driver_name,

                "driver_mobile": partner.driver_mobile,

                "vehicle_number": partner.vehicle_number,

                "cab_category": partner.cab_category,
            }

            if partner

            else None
        ),
    }

    return Response(data)


# =========================================================
# EXPORT BOOKING REPORT
# =========================================================

@api_view(["GET"])
def export_booking_report(request):

    bookings = Booking.objects.select_related(
        "partner"
    ).all().order_by(
        "-created_at"
    )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Booking Report"

    headers = [
        "Booking ID",
        "Customer Name",
        "Mobile",
        "Service Type",
        "Package",
        "Travel Date",
        "Passengers",
        "Cab Category",
        "Pickup",
        "Drop",
        "Amount",
        "Payment Status",
        "Booking Status",
        "Partner",
        "Created At",
    ]

    sheet.append(headers)

    for booking in bookings:

        sheet.append([

            booking.booking_id,

            booking.customer_name,

            booking.mobile,

            booking.service_type,

            booking.package_name,

            booking.travel_date,

            booking.passengers,

            booking.cab_category,

            booking.pickup_location,

            booking.drop_location,

            float(booking.amount),

            booking.payment_status,

            booking.booking_status,

            (
                booking.partner.partner_name
                if booking.partner
                else "Unassigned"
            ),

            booking.created_at.strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
        ])

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="SR_Journey_Booking_Report.xlsx"'
    )

    workbook.save(response)

    return response


# =========================================================
# EXPORT REVENUE REPORT
# =========================================================

@api_view(["GET"])
def export_revenue_report(request):

    bookings = Booking.objects.filter(
        booking_status__in=[
            "CONFIRMED",
            "COMPLETED"
        ]
    ).order_by(
        "-created_at"
    )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Revenue Report"

    headers = [
        "Booking ID",
        "Customer Name",
        "Service Type",
        "Package",
        "Travel Date",
        "Amount",
        "Payment Status",
        "Booking Status",
        "Created At",
    ]

    sheet.append(headers)

    for booking in bookings:

        sheet.append([

            booking.booking_id,

            booking.customer_name,

            booking.service_type,

            booking.package_name,

            booking.travel_date,

            float(booking.amount),

            booking.payment_status,

            booking.booking_status,

            booking.created_at.strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
        ])

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="SR_Journey_Revenue_Report.xlsx"'
    )

    workbook.save(response)

    return response


# =========================================================
# EXPORT PARTNER REPORT
# =========================================================

@api_view(["GET"])
def export_partner_report(request):

    partners = Partner.objects.all().order_by(
        "partner_name"
    )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Partner Report"

    headers = [
        "Partner Name",
        "Company",
        "Mobile",
        "Cab Category",
        "Vehicle Number",
        "Driver Name",
        "Driver Mobile",
        "Status",
        "Total Bookings",
        "Total Revenue",
    ]

    sheet.append(headers)

    for partner in partners:

        bookings = Booking.objects.filter(
            partner=partner
        )

        total_bookings = bookings.count()

        total_revenue = sum(

            float(booking.amount)

            for booking in bookings

            if booking.booking_status
            in [
                "CONFIRMED",
                "COMPLETED"
            ]
        )

        sheet.append([

            partner.partner_name,

            partner.company_name,

            partner.mobile,

            partner.cab_category,

            partner.vehicle_number,

            partner.driver_name,

            partner.driver_mobile,

            partner.status,

            total_bookings,

            total_revenue,
        ])

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="SR_Journey_Partner_Report.xlsx"'
    )

    workbook.save(response)

    return response


# =========================================================
# EXPORT COMMISSION REPORT
# =========================================================

@api_view(["GET"])
def export_commission_report(request):

    bookings = Booking.objects.select_related(
        "partner"
    ).order_by("-created_at")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Commission Report"

    headers = [
        "Booking ID",
        "Customer Name",
        "Mobile",
        "Partner Name",
        "Booking Amount",
        "Commission Rate (%)",
        "Commission Amount",
        "Partner Payable",
        "Booking Status",
        "Payment Status",
        "Travel Date",
        "Created At",
    ]

    sheet.append(headers)

    for booking in bookings:
        sheet.append([
            booking.booking_id,
            booking.customer_name,
            booking.mobile,
            booking.partner.partner_name if booking.partner else "Unassigned",
            float(booking.amount or 0),
            float(booking.commission_rate or 0),
            float(booking.commission_amount or 0),
            float(booking.partner_payable or 0),
            booking.booking_status,
            booking.payment_status,
            str(booking.travel_date or ""),
            str(booking.created_at),
        ])

    # Format header row
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)

    # Set readable column widths
    for column_cells in sheet.columns:
        column_letter = column_cells[0].column_letter
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )
        sheet.column_dimensions[column_letter].width = min(
            max_length + 2, 30
        )

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="SR_Journey_Commission_Report.xlsx"'
    )

    workbook.save(response)
    return response


# =========================================================
# BOOKING DETAIL PAGE
# =========================================================

@staff_member_required
def booking_detail(request, booking_id):

    booking = get_object_or_404(
        Booking.objects.select_related("partner"),
        booking_id=booking_id
    )

    partners = Partner.objects.filter(
        status="ACTIVE"
    ).order_by("partner_name")

    return render(
        request,
        "travelbookings/booking_detail.html",
        {
            "booking": booking,
            "partners": partners,
        }
    )


# =========================================================
# BOOKING ACTION
# CONFIRM / CANCEL / COMPLETE
# =========================================================

@staff_member_required
@require_POST
def booking_action(request, booking_id, action):

    booking = get_object_or_404(
        Booking.objects.select_related(
            "partner"
        ),
        booking_id=booking_id
    )

    if action == "confirm":

        booking.booking_status = "CONFIRMED"

        booking.save(
            update_fields=[
                "booking_status"
            ]
        )

    elif action == "cancel":

        booking.booking_status = "CANCELLED"

        booking.save(
            update_fields=[
                "booking_status"
            ]
        )

    elif action == "complete":

        booking.booking_status = "COMPLETED"

        booking.save(
            update_fields=[
                "booking_status"
            ]
        )

    return redirect(
        "booking_detail",
        booking_id=booking.booking_id
    )
    # =========================================================
# PAYMENT ACTION
# PAID / PENDING / FAILED / REFUNDED
# =========================================================

@staff_member_required
@require_POST
def payment_action(request, booking_id, payment_status):

    booking = get_object_or_404(
        Booking.objects.select_related(
            "partner"
        ),
        booking_id=booking_id
    )

    allowed_statuses = [
        "PENDING",
        "PAID",
        "FAILED",
        "REFUNDED",
    ]

    if payment_status in allowed_statuses:

        booking.payment_status = payment_status

        booking.save(
            update_fields=[
                "payment_status"
            ]
        )

    return redirect(
        "booking_detail",
        booking_id=booking.booking_id
    )
    # =========================================================
# PARTNER ASSIGN / CHANGE
# =========================================================


from decimal import Decimal, ROUND_HALF_UP

@staff_member_required
@require_POST
def assign_partner(request, booking_id):
    booking = get_object_or_404(
        Booking.objects.select_related("partner"),
        booking_id=booking_id
    )

    partner_id = request.POST.get("partner_id")

    if partner_id:
        partner = get_object_or_404(
            Partner,
            id=partner_id,
            status="ACTIVE"
        )
        booking.partner = partner

        # Calculate 5% commission
        amount = booking.amount or Decimal("0.00")
        rate = booking.commission_rate or Decimal("5.00")

        commission = (
            amount * rate / Decimal("100")
        ).quantize(
            Decimal("0.01"),
            rounding=ROUND_HALF_UP
        )

        booking.commission_amount = commission
        booking.partner_payable = amount - commission

        booking.save(update_fields=[
            "partner",
            "commission_rate",
            "commission_amount",
            "partner_payable",
        ])

    else:
        # Partner removed: no partner payable
        booking.partner = None
        booking.partner_payable = Decimal("0.00")

        booking.save(update_fields=[
            "partner",
            "partner_payable",
        ])

    return redirect(
        "booking_detail",
        booking_id=booking.booking_id
    )

@staff_member_required
def partner_booking_history(request, partner_id):
    partner = get_object_or_404(Partner, id=partner_id)

    bookings = Booking.objects.filter(
        partner=partner
    ).order_by("-created_at")

    summary = bookings.aggregate(
        total_revenue=Sum("amount")
    )

    context = {
        "partner": partner,
        "bookings": bookings,
        "total_bookings": bookings.count(),
        "completed_bookings": bookings.filter(
            booking_status="COMPLETED"
        ).count(),
        "pending_bookings": bookings.filter(
            booking_status="PENDING"
        ).count(),
        "total_revenue": summary["total_revenue"] or 0,
    }

    return render(
        request,
        "travelbookings/partner_booking_history.html",
        context,
    )
    
# =========================================================
# PARTNER SETTLEMENT RECEIPT
# =========================================================

@staff_member_required
def settlement_receipt(request, settlement_id):
    settlement = get_object_or_404(
        PartnerSettlement.objects.select_related("partner"),
        id=settlement_id
    )

    total_payable = Booking.objects.filter(
        partner=settlement.partner,
        booking_status__in=["CONFIRMED", "COMPLETED"],
    ).aggregate(
        total=Sum("partner_payable")
    )["total"] or 0

    total_paid = PartnerSettlement.objects.filter(
        partner=settlement.partner,
        payment_date__lte=settlement.payment_date,
    ).aggregate(
        total=Sum("amount")
    )["total"] or 0

    pending_balance = total_payable - total_paid

    context = {
        "settlement": settlement,
        "partner": settlement.partner,
        "total_payable": total_payable,
        "total_paid": total_paid,
        "pending_balance": pending_balance,
    }

    return render(
        request,
        "travelbookings/settlement_receipt.html",
        context,
    )


# =========================================================
# PARTNER SETTLEMENT EXCEL REPORT
# =========================================================

@staff_member_required
def export_settlement_report(request):

    settlements = PartnerSettlement.objects.select_related(
        "partner"
    ).order_by("-payment_date", "-id")

    # Read the same filters used on settlement history
    partner_id = request.GET.get("partner", "").strip()
    start_date = request.GET.get("start_date", "").strip()
    end_date = request.GET.get("end_date", "").strip()

    # Partner filter
    if partner_id.isdigit():
        settlements = settlements.filter(
            partner_id=int(partner_id)
        )

    # Start date filter
    if start_date:
        try:
            datetime.strptime(start_date, "%Y-%m-%d")
            settlements = settlements.filter(
                payment_date__gte=start_date
            )
        except ValueError:
            pass

    # End date filter
    if end_date:
        try:
            datetime.strptime(end_date, "%Y-%m-%d")
            settlements = settlements.filter(
                payment_date__lte=end_date
            )
        except ValueError:
            pass

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Settlement Report"

    headers = [
        "Settlement ID",
        "Partner Name",
        "Vehicle Number",
        "Amount Paid",
        "Payment Date",
        "Payment Method",
        "Transaction Reference",
        "Notes",
    ]

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = cell.font.copy(bold=True)

    for settlement in settlements:
        worksheet.append([
            settlement.id,
            settlement.partner.partner_name,
            settlement.partner.vehicle_number,
            float(settlement.amount),
            settlement.payment_date.strftime("%Y-%m-%d"),
            settlement.get_payment_method_display(),
            settlement.transaction_reference or "",
            settlement.notes or "",
        ])

    column_widths = {
        "A": 16,
        "B": 25,
        "C": 20,
        "D": 18,
        "E": 18,
        "F": 20,
        "G": 28,
        "H": 35,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="SR_Journey_Settlement_Report.xlsx"'
    )

    workbook.save(response)

    return response


# =========================================================
# PARTNER SETTLEMENT HISTORY
# =========================================================

@staff_member_required
def settlement_history(request):
    settlements = PartnerSettlement.objects.select_related(
        "partner"
    ).order_by("-payment_date", "-id")

    # Read filters
    partner_id = request.GET.get("partner", "").strip()
    start_date = request.GET.get("start_date", "").strip()
    end_date = request.GET.get("end_date", "").strip()

    # Partner filter
    if partner_id.isdigit():
        settlements = settlements.filter(partner_id=int(partner_id))
    else:
        partner_id = ""

    # Start date filter
    if start_date:
        try:
            datetime.strptime(start_date, "%Y-%m-%d")
            settlements = settlements.filter(
                payment_date__gte=start_date
            )
        except ValueError:
            start_date = ""

    # End date filter
    if end_date:
        try:
            datetime.strptime(end_date, "%Y-%m-%d")
            settlements = settlements.filter(
                payment_date__lte=end_date
            )
        except ValueError:
            end_date = ""

    # Filtered summary
    total_settlements = settlements.count()

    total_paid = settlements.aggregate(
        total=Sum("amount")
    )["total"] or 0

    context = {
        "settlements": settlements,
        "total_settlements": total_settlements,
        "total_paid": total_paid,
        "partners": Partner.objects.all().order_by("partner_name"),
        "selected_partner": partner_id,
        "start_date": start_date,
        "end_date": end_date,
    }

    return render(
        request,
        "travelbookings/settlement_history.html",
        context,
    )
    
    
# =========================================================
# PARTNER PAYMENT DASHBOARD
# =========================================================

@staff_member_required
def partner_payment_dashboard(request):

    partners = Partner.objects.all().order_by("partner_name")

    partner_data = []

    grand_total_payable = 0
    grand_total_paid = 0
    grand_total_pending = 0

    for partner in partners:

        # Payable from confirmed/completed bookings
        total_payable = Booking.objects.filter(
            partner=partner,
            booking_status__in=["CONFIRMED", "COMPLETED"],
        ).aggregate(
            total=Sum("partner_payable")
        )["total"] or 0

        # All settlement payments recorded for this partner
        total_paid = PartnerSettlement.objects.filter(
            partner=partner
        ).aggregate(
            total=Sum("amount")
        )["total"] or 0

        pending_balance = total_payable - total_paid

        partner_data.append({
            "partner": partner,
            "total_payable": total_payable,
            "total_paid": total_paid,
            "pending_balance": pending_balance,
        })

        grand_total_payable += total_payable
        grand_total_paid += total_paid
        grand_total_pending += pending_balance

    context = {
        "partner_data": partner_data,
        "grand_total_payable": grand_total_payable,
        "grand_total_paid": grand_total_paid,
        "grand_total_pending": grand_total_pending,
    }

    return render(
        request,
        "travelbookings/partner_payment_dashboard.html",
        context,
    )



